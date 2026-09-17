import os
import json
import openpyxl

clss_dir = r"c:\Users\Ashish\OneDrive - Absolute Return For Kids\CLSS RF"
excel_path = os.path.join(clss_dir, "CLSS_ToC & RF_26-27.xlsx")

wb = openpyxl.load_workbook(excel_path, data_only=True)

# Parse ToC
toc_ws = wb["ToC_CLSS_26-27"]
toc_items = []
for r in range(4, min(30, toc_ws.max_row + 1)):
    outcome = toc_ws.cell(r, 2).value
    indicators = toc_ws.cell(r, 3).value
    content = toc_ws.cell(r, 4).value
    pmu = toc_ws.cell(r, 5).value
    state = toc_ws.cell(r, 6).value
    mel = toc_ws.cell(r, 7).value
    if outcome or indicators:
        toc_items.append({
            "outcome": str(outcome or "").strip(),
            "indicators": str(indicators or "").strip(),
            "content_actions": str(content or "").strip(),
            "pmu_actions": str(pmu or "").strip(),
            "state_actions": str(state or "").strip(),
            "mel_actions": str(mel or "").strip()
        })

# Parse Result Framework
rf_ws = wb["Result Framework (3)"]
rf_indicators = []
for r in range(4, rf_ws.max_row + 1):
    sn = rf_ws.cell(r, 1).value
    level = rf_ws.cell(r, 2).value
    statement = rf_ws.cell(r, 3).value
    ind_text = rf_ws.cell(r, 4).value
    tool = rf_ws.cell(r, 9).value if rf_ws.max_column >= 9 else ""
    freq = rf_ws.cell(r, 11).value if rf_ws.max_column >= 11 else ""
    target = rf_ws.cell(r, 5).value if rf_ws.max_column >= 5 else ""
    ind_type = rf_ws.cell(r, 8).value if rf_ws.max_column >= 8 else ""
    if statement or ind_text:
        rf_indicators.append({
            "sn": str(sn or "").strip(),
            "level": str(level or "").strip(),
            "statement": str(statement or "").strip(),
            "indicator": str(ind_text or "").strip(),
            "tool": str(tool or "").strip(),
            "frequency": str(freq or "").strip(),
            "target": str(target or "").strip(),
            "type": str(ind_type or "").strip()
        })

# Parse Definitions
def_ws = wb["Definition sheet 26-27"]
definitions = []
for r in range(2, def_ws.max_row + 1):
    sn = def_ws.cell(r, 1).value
    ind = def_ws.cell(r, 2).value
    term = def_ws.cell(r, 3).value
    defn = def_ws.cell(r, 4).value
    rem = def_ws.cell(r, 5).value
    if term or defn:
        definitions.append({
            "sn": sn,
            "indicator": str(ind or "").strip(),
            "term": str(term or "").strip(),
            "definition": str(defn or "").strip(),
            "remark": str(rem or "").strip()
        })

toc_json = json.dumps(toc_items, ensure_ascii=False)
rf_json = json.dumps(rf_indicators, ensure_ascii=False)
def_json = json.dumps(definitions, ensure_ascii=False)

html_template = f"""<!DOCTYPE html>
<html lang="en">
<head>
  <meta charset="UTF-8">
  <meta name="viewport" content="width=device-width, initial-scale=1.0">
  <title>CLSS Results Framework & Theory of Change Hub (AY 2026-27)</title>
  <link rel="preconnect" href="https://fonts.googleapis.com">
  <link rel="preconnect" href="https://fonts.gstatic.com" crossorigin>
  <link href="https://fonts.googleapis.com/css2?family=Inter:wght@300;400;500;600;700;800&family=JetBrains+Mono:wght@400;500;600&display=swap" rel="stylesheet">
  <style>
    :root {{
      --bg-main: #f8fafc;
      --bg-card: #ffffff;
      --bg-subtle: #f1f5f9;
      --text-main: #0f172a;
      --text-muted: #64748b;
      --primary: #4338ca;
      --primary-light: #e0e7ff;
      --primary-hover: #3730a3;
      --accent-teal: #0d9488;
      --accent-amber: #d97706;
      --accent-purple: #7e22ce;
      --accent-emerald: #059669;
      --border: #e2e8f0;
      --border-focus: #6366f1;
      --shadow-sm: 0 1px 2px 0 rgb(0 0 0 / 0.05);
      --shadow-md: 0 4px 6px -1px rgb(0 0 0 / 0.1), 0 2px 4px -2px rgb(0 0 0 / 0.1);
      --radius: 12px;
    }}

    * {{ box-sizing: border-box; margin: 0; padding: 0; }}
    body {{
      font-family: 'Inter', -apple-system, BlinkMacSystemFont, sans-serif;
      background-color: var(--bg-main);
      color: var(--text-main);
      line-height: 1.5;
      -webkit-font-smoothing: antialiased;
    }}

    .top-header {{
      background: linear-gradient(135deg, #1e1b4b 0%, #312e81 50%, #4338ca 100%);
      color: #ffffff;
      padding: 2.5rem 2rem 2rem;
      position: relative;
      overflow: hidden;
    }}
    .header-container {{
      max-width: 1300px;
      margin: 0 auto;
    }}
    .badge {{
      display: inline-flex;
      align-items: center;
      gap: 6px;
      background: rgba(255, 255, 255, 0.15);
      backdrop-filter: blur(8px);
      padding: 4px 12px;
      border-radius: 9999px;
      font-size: 0.75rem;
      font-weight: 600;
      letter-spacing: 0.05em;
      text-transform: uppercase;
      margin-bottom: 0.75rem;
      border: 1px solid rgba(255, 255, 255, 0.2);
    }}
    .title-row {{
      display: flex;
      justify-content: space-between;
      align-items: flex-start;
      flex-wrap: wrap;
      gap: 1.5rem;
    }}
    .top-header h1 {{
      font-size: 2rem;
      font-weight: 800;
      letter-spacing: -0.02em;
      margin-bottom: 0.5rem;
    }}
    .top-header p {{
      color: #c7d2fe;
      font-size: 0.95rem;
      max-width: 750px;
    }}
    .header-links {{
      display: flex;
      gap: 0.75rem;
      margin-top: 0.5rem;
    }}
    .btn-pill {{
      background: rgba(255, 255, 255, 0.12);
      color: white;
      text-decoration: none;
      padding: 8px 16px;
      border-radius: 8px;
      font-size: 0.85rem;
      font-weight: 600;
      border: 1px solid rgba(255, 255, 255, 0.25);
    }}

    .main-container {{
      max-width: 1300px;
      margin: -1.5rem auto 3rem;
      padding: 0 1.5rem;
    }}

    .kpi-grid {{
      display: grid;
      grid-template-columns: repeat(auto-fit, minmax(180px, 1fr));
      gap: 1rem;
      margin-bottom: 2rem;
    }}
    .kpi-card {{
      background: var(--bg-card);
      border: 1px solid var(--border);
      border-radius: var(--radius);
      padding: 1.25rem;
      box-shadow: var(--shadow-sm);
      position: relative;
    }}
    .kpi-card::before {{
      content: '';
      position: absolute;
      top: 0;
      left: 0;
      width: 4px;
      height: 100%;
      background: var(--primary);
    }}
    .kpi-card.accent-teal::before {{ background: var(--accent-teal); }}
    .kpi-card.accent-amber::before {{ background: var(--accent-amber); }}
    .kpi-card.accent-purple::before {{ background: var(--accent-purple); }}
    .kpi-card.accent-emerald::before {{ background: var(--accent-emerald); }}

    .kpi-label {{
      font-size: 0.75rem;
      font-weight: 600;
      color: var(--text-muted);
      text-transform: uppercase;
      letter-spacing: 0.05em;
    }}
    .kpi-value {{
      font-size: 1.75rem;
      font-weight: 800;
      color: var(--text-main);
      line-height: 1.2;
      margin: 4px 0;
    }}
    .kpi-subtext {{
      font-size: 0.75rem;
      color: var(--text-muted);
    }}

    .tabs-nav {{
      display: flex;
      gap: 0.5rem;
      border-bottom: 1px solid var(--border);
      margin-bottom: 1.5rem;
      background: var(--bg-card);
      padding: 0.5rem 0.5rem 0;
      border-radius: var(--radius) var(--radius) 0 0;
      box-shadow: var(--shadow-sm);
      flex-wrap: wrap;
    }}
    .tab-btn {{
      padding: 0.75rem 1.25rem;
      background: none;
      border: none;
      border-bottom: 3px solid transparent;
      font-size: 0.9rem;
      font-weight: 600;
      color: var(--text-muted);
      cursor: pointer;
      display: flex;
      align-items: center;
      gap: 8px;
    }}
    .tab-btn:hover {{
      color: var(--primary);
    }}
    .tab-btn.active {{
      color: var(--primary);
      border-bottom-color: var(--primary);
      background: var(--primary-light);
      border-radius: 8px 8px 0 0;
    }}

    .tab-content {{
      display: none;
    }}
    .tab-content.active {{
      display: block;
    }}

    .section-card {{
      background: var(--bg-card);
      border: 1px solid var(--border);
      border-radius: var(--radius);
      padding: 1.75rem;
      box-shadow: var(--shadow-sm);
      margin-bottom: 1.5rem;
    }}
    .section-title {{
      font-size: 1.25rem;
      font-weight: 700;
      color: var(--text-main);
      margin-bottom: 0.25rem;
    }}
    .section-subtitle {{
      font-size: 0.875rem;
      color: var(--text-muted);
      margin-bottom: 1.5rem;
    }}

    .toc-grid {{
      display: grid;
      grid-template-columns: repeat(auto-fit, minmax(360px, 1fr));
      gap: 1.25rem;
    }}
    .toc-card {{
      background: #ffffff;
      border: 1px solid var(--border);
      border-radius: 10px;
      padding: 1.25rem;
      box-shadow: var(--shadow-sm);
      border-left: 4px solid var(--primary);
    }}
    .toc-header {{
      display: flex;
      align-items: flex-start;
      gap: 0.75rem;
      margin-bottom: 0.75rem;
    }}
    .toc-number {{
      background: var(--primary-light);
      color: var(--primary);
      font-weight: 800;
      font-size: 0.85rem;
      width: 28px;
      height: 28px;
      border-radius: 6px;
      display: flex;
      align-items: center;
      justify-content: center;
      flex-shrink: 0;
    }}
    .toc-title {{
      font-size: 1rem;
      font-weight: 700;
      color: var(--text-main);
    }}
    .toc-indicators {{
      font-size: 0.825rem;
      color: #334155;
      background: var(--bg-subtle);
      padding: 0.75rem;
      border-radius: 6px;
      margin-bottom: 0.75rem;
      line-height: 1.45;
      white-space: pre-line;
    }}
    .raci-tags {{
      display: flex;
      flex-wrap: wrap;
      gap: 0.35rem;
      font-size: 0.75rem;
    }}
    .raci-pill {{
      padding: 2px 8px;
      border-radius: 4px;
      font-weight: 600;
    }}
    .pill-content {{ background: #dbeafe; color: #1e40af; }}
    .pill-pmu {{ background: #fef3c7; color: #92400e; }}
    .pill-state {{ background: #e0e7ff; color: #3730a3; }}
    .pill-mel {{ background: #ccfbf1; color: #115e59; }}

    .filter-bar {{
      display: flex;
      gap: 1rem;
      margin-bottom: 1.25rem;
      flex-wrap: wrap;
      align-items: center;
    }}
    .search-input {{
      flex: 1;
      min-width: 260px;
      padding: 0.65rem 1rem;
      border: 1px solid var(--border);
      border-radius: 8px;
      font-size: 0.875rem;
      outline: none;
    }}
    .search-input:focus {{
      border-color: var(--border-focus);
    }}
    .filter-select {{
      padding: 0.65rem 1rem;
      border: 1px solid var(--border);
      border-radius: 8px;
      font-size: 0.875rem;
      background: white;
      cursor: pointer;
      outline: none;
    }}

    .table-container {{
      overflow-x: auto;
      border: 1px solid var(--border);
      border-radius: 8px;
    }}
    table {{
      width: 100%;
      border-collapse: collapse;
      font-size: 0.85rem;
      text-align: left;
    }}
    th {{
      background: #f8fafc;
      color: var(--text-muted);
      font-weight: 700;
      padding: 0.75rem 1rem;
      border-bottom: 1px solid var(--border);
      text-transform: uppercase;
      font-size: 0.75rem;
      letter-spacing: 0.05em;
    }}
    td {{
      padding: 0.85rem 1rem;
      border-bottom: 1px solid var(--border);
      vertical-align: top;
      color: var(--text-main);
    }}
    tr:hover td {{ background: #f8fafc; }}

    .tag-level {{
      display: inline-block;
      padding: 2px 6px;
      border-radius: 4px;
      font-size: 0.7rem;
      font-weight: 700;
      text-transform: uppercase;
    }}
    .level-outcome {{ background: #ede9fe; color: #6d28d9; }}
    .level-output {{ background: #dcfce7; color: #15803d; }}
    .level-impact {{ background: #fee2e2; color: #b91c1c; }}
    .level-intermediate {{ background: #e0f2fe; color: #0369a1; }}

    .rubric-card {{
      border: 1px solid var(--border);
      border-radius: 8px;
      padding: 1.25rem;
      margin-bottom: 1rem;
      background: #ffffff;
    }}
    .rubric-term {{
      font-weight: 700;
      color: var(--primary);
      font-size: 1.05rem;
      margin-bottom: 0.25rem;
    }}
    .qualifier-box {{
      background: #f0fdf4;
      border: 1px solid #bbf7d0;
      border-radius: 6px;
      padding: 0.75rem 1rem;
      margin: 0.75rem 0;
      font-size: 0.85rem;
      color: #166534;
    }}
  </style>
</head>
<body>

  <header class="top-header">
    <div class="header-container">
      <div class="badge">State Program Dashboard • Academic Year 2026–27</div>
      <div class="title-row">
        <div>
          <h1>CLSS Results Framework & Theory of Change</h1>
          <p>Cluster-Level Shaikshik Samvad (CLSS) MEL Infrastructure, Indicator Matrix, Scale Metrics, and Operational Guidelines for Madhya Pradesh.</p>
        </div>
        <div class="header-links">
          <span class="btn-pill">📊 47 Indicators</span>
          <span class="btn-pill">🏛️ 55 Districts</span>
        </div>
      </div>
    </div>
  </header>

  <main class="main-container">

    <div class="kpi-grid">
      <div class="kpi-card">
        <div class="kpi-label">Districts</div>
        <div class="kpi-value">55</div>
        <div class="kpi-subtext">District Core Committees (DCC)</div>
      </div>
      <div class="kpi-card accent-teal">
        <div class="kpi-label">Blocks</div>
        <div class="kpi-value">~322</div>
        <div class="kpi-subtext">Sub-district units</div>
      </div>
      <div class="kpi-card accent-purple">
        <div class="kpi-label">Clusters & CACs</div>
        <div class="kpi-value">3,066</div>
        <div class="kpi-subtext">Cluster facilitators</div>
      </div>
      <div class="kpi-card accent-amber">
        <div class="kpi-label">CLSS Cycles</div>
        <div class="kpi-value">5</div>
        <div class="kpi-subtext">Thematic Margdarshikas</div>
      </div>
      <div class="kpi-card accent-emerald">
        <div class="kpi-label">Evaluation Tools</div>
        <div class="kpi-value">4</div>
        <div class="kpi-subtext">Obs, Feedback, Attendance</div>
      </div>
      <div class="kpi-card">
        <div class="kpi-label">ToC Outcomes</div>
        <div class="kpi-value">9</div>
        <div class="kpi-subtext">Core impact streams</div>
      </div>
    </div>

    <div class="tabs-nav">
      <button class="tab-btn active" id="btn-toc" onclick="switchTab('toc')">🎯 Theory of Change (9 Outcomes)</button>
      <button class="tab-btn" id="btn-rf" onclick="switchTab('rf')">📊 Results Framework Explorer (47 Indicators)</button>
      <button class="tab-btn" id="btn-definitions" onclick="switchTab('definitions')">📐 Rubrics & Standard Definitions</button>
      <button class="tab-btn" id="btn-guidelines" onclick="switchTab('guidelines')">📋 8-Step MEL Governance</button>
    </div>

    <div id="tab-toc" class="tab-content active">
      <div class="section-card">
        <div class="section-title">🎯 Theory of Change (ToC) Framework</div>
        <div class="section-subtitle">The 9 foundational outcome streams driving teacher peer learning, pedagogical adoption, and system institutionalization.</div>
        <div class="toc-grid" id="toc-container"></div>
      </div>
    </div>

    <div id="tab-rf" class="tab-content">
      <div class="section-card">
        <div class="section-title">📊 CLSS Results Framework Indicator Matrix</div>
        <div class="section-subtitle">Detailed indicators, levels, target thresholds, data collection tools, and reporting frequencies.</div>

        <div class="filter-bar">
          <input type="text" id="rf-search" class="search-input" placeholder="🔍 Search indicators, tools, or statements..." oninput="filterIndicators()">
          <select id="rf-level-filter" class="filter-select" onchange="filterIndicators()">
            <option value="">All Levels</option>
            <option value="Outcome">Outcome</option>
            <option value="Output">Output</option>
            <option value="Intermediate Outcome">Intermediate Outcome</option>
            <option value="Impact">Impact</option>
          </select>
          <select id="rf-type-filter" class="filter-select" onchange="filterIndicators()">
            <option value="">All Categories</option>
            <option value="Knowledge">Knowledge</option>
            <option value="Skill">Skill</option>
            <option value="Mindset">Mindset</option>
            <option value="Process">Process</option>
          </select>
        </div>

        <div class="table-container">
          <table>
            <thead>
              <tr>
                <th style="width: 60px;">S.N.</th>
                <th style="width: 140px;">Level</th>
                <th>Statement / Objective</th>
                <th>Indicator Description</th>
                <th style="width: 140px;">Tool</th>
                <th style="width: 100px;">Frequency</th>
                <th style="width: 90px;">Category</th>
              </tr>
            </thead>
            <tbody id="rf-table-body"></tbody>
          </table>
        </div>
      </div>
    </div>

    <div id="tab-definitions" class="tab-content">
      <div class="section-card">
        <div class="section-title">📐 Standard Definitions & Measurement Rubrics</div>
        <div class="section-subtitle">Locked rubrics aligned across Content, PMU, State, and MEL leadership.</div>

        <div class="rubric-card">
          <div class="rubric-term">1. Definition of 'Best Practice' (Indicator 14.4)</div>
          <p>During CLSS, a <strong>Best Practice</strong> is counted only when a teacher shares an authentic classroom experience describing how they used a Shaikshik Samwad technique to solve a specific challenge, resulting in positive student change.</p>
          <div class="qualifier-box">
            <strong>Required 3-Part Qualifiers (Must all be met by a single teacher):</strong><br>
            1. <strong>Classroom Challenge:</strong> Explicitly states the specific pedagogical challenge encountered.<br>
            2. <strong>Samwad Learning Applied:</strong> Details the exact CLSS learning/technique utilized.<br>
            3. <strong>Observable Student Impact:</strong> Shares the result as a positive change in student learning or engagement.
          </div>
          <p><em>Example:</em> "Students struggled with thinking time, so I introduced a '30-second silent think hand-signal' before cold calling, which doubled response participation."</p>
        </div>

        <div class="rubric-card">
          <div class="rubric-term">2. Participant Engagement (>=50% Threshold)</div>
          <p>Defined as the proportion of participating teachers who actively contribute individual reflections, ask/answer questions during structured discussions, and submit valid session feedback forms.</p>
        </div>

        <div class="rubric-card">
          <div class="rubric-term">3. Facilitator Proficiency Scale (Beginner to Expert)</div>
          <p>Observation rubric tracking whether facilitators utilize the 7 design principles, explain learning objectives clearly, maintain time-boxed Margdarshika flow, and cultivate psychological safety.</p>
        </div>
      </div>
    </div>

    <div id="tab-guidelines" class="tab-content">
      <div class="section-card">
        <div class="section-title">📋 8-Step MEL Governance & Workflow Protocol</div>
        <div class="section-subtitle">Standard operating procedure for operationalizing the Results Framework.</div>

        <div style="display: grid; gap: 1rem;">
          <div class="rubric-card"><strong>Step 1: Theory of Change Alignment</strong> — Program & MEL teams co-create and lock the ToC template.</div>
          <div class="rubric-card"><strong>Step 2: Indicator Formulation</strong> — Map outputs, intermediate outcomes, and impact statements to quantifiable indicators.</div>
          <div class="rubric-card"><strong>Step 3: Tool Template Design</strong> — MEL drafts standard observation, feedback, attendance, and facilitator forms.</div>
          <div class="rubric-card"><strong>Step 4: Question Framing</strong> — Content & Program teams frame contextual questions following Question Framing Guidelines.</div>
          <div class="rubric-card"><strong>Step 5: Review & Finalization</strong> — Joint working review to finalize tool questionnaires.</div>
          <div class="rubric-card"><strong>Step 6: Field Piloting</strong> — Pilot test each tool prior to statewide deployment.</div>
          <div class="rubric-card"><strong>Step 7: Ongoing Monitoring</strong> — Real-time tracking of indicator submissions across clusters.</div>
          <div class="rubric-card"><strong>Step 8: Target Achievement & Iteration</strong> — Periodic analysis briefs and review meetings to adjust strategy.</div>
        </div>
      </div>
    </div>

  </main>

  <script>
    const tocData = {toc_json};
    const rfData = {rf_json};

    function switchTab(tabId) {{
      document.querySelectorAll('.tab-btn').forEach(btn => btn.classList.remove('active'));
      document.querySelectorAll('.tab-content').forEach(content => content.classList.remove('active'));
      
      const targetBtn = document.getElementById('btn-' + tabId);
      if (targetBtn) targetBtn.classList.add('active');
      
      const targetContent = document.getElementById('tab-' + tabId);
      if (targetContent) targetContent.classList.add('active');
    }}

    function renderToC() {{
      const container = document.getElementById('toc-container');
      container.innerHTML = tocData.map((item, idx) => {{
        const cp = item.content_actions ? '<span class="raci-pill pill-content">Content: ' + item.content_actions.substring(0, 45) + '...</span>' : '';
        const pp = item.pmu_actions ? '<span class="raci-pill pill-pmu">PMU: ' + item.pmu_actions.substring(0, 45) + '...</span>' : '';
        const mp = item.mel_actions ? '<span class="raci-pill pill-mel">MEL: ' + item.mel_actions.substring(0, 45) + '...</span>' : '';
        return '<div class="toc-card"><div class="toc-header"><div class="toc-number">' + (idx + 1) + '</div><div class="toc-title">' + item.outcome + '</div></div><div class="toc-indicators">' + (item.indicators || 'Core outcome stream') + '</div><div class="raci-tags">' + cp + pp + mp + '</div></div>';
      }}).join('');
    }}

    function renderRF(data) {{
      const tbody = document.getElementById('rf-table-body');
      tbody.innerHTML = data.map(item => {{
        let levelClass = 'level-outcome';
        const lvl = (item.level || '').toLowerCase();
        if (lvl.includes('output')) levelClass = 'level-output';
        else if (lvl.includes('impact')) levelClass = 'level-impact';
        else if (lvl.includes('intermediate')) levelClass = 'level-intermediate';

        return '<tr><td><strong>' + (item.sn || '-') + '</strong></td><td><span class="tag-level ' + levelClass + '">' + (item.level || 'Outcome') + '</span></td><td><strong>' + (item.statement || '-') + '</strong></td><td>' + (item.indicator || '-') + '</td><td><code style="background:#f1f5f9; padding:2px 6px; border-radius:4px;">' + (item.tool || 'Document') + '</code></td><td>' + (item.frequency || 'Monthly') + '</td><td><span style="font-size:0.75rem; font-weight:600; color:#475569;">' + (item.type || 'Skill') + '</span></td></tr>';
      }}).join('');
    }}

    function filterIndicators() {{
      const query = (document.getElementById('rf-search').value || '').toLowerCase();
      const levelFilter = (document.getElementById('rf-level-filter').value || '').toLowerCase();
      const typeFilter = (document.getElementById('rf-type-filter').value || '').toLowerCase();

      const filtered = rfData.filter(item => {{
        const matchesQuery = !query || 
          (item.statement || '').toLowerCase().includes(query) || 
          (item.indicator || '').toLowerCase().includes(query) || 
          (item.tool || '').toLowerCase().includes(query);
        const matchesLevel = !levelFilter || (item.level || '').toLowerCase().includes(levelFilter);
        const matchesType = !typeFilter || (item.type || '').toLowerCase().includes(typeFilter);
        return matchesQuery && matchesLevel && matchesType;
      }});

      renderRF(filtered);
    }}

    renderToC();
    renderRF(rfData);
  </script>
</body>
</html>
"""

dash_path = os.path.join(clss_dir, "clss_rf_dashboard.html")
with open(dash_path, "w", encoding="utf-8") as f:
    f.write(html_template)

print(f"CLSS Dashboard successfully created at: {dash_path}")
